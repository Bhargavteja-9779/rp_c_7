import sys, docx
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = docx.Document(sys.argv[1])
sec = d.sections[0]; sec.page_width = Cm(21.0); sec.page_height = Cm(29.7); sec.left_margin = sec.right_margin = Cm(2.0); usable = sec.page_width - sec.left_margin - sec.right_margin
for t in d.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    n = len(t.columns); lens = [max(len(c.text) for c in col.cells) for col in t.columns]
    w = [max(7, min(l, 40)) for l in lens]; tot = sum(w)
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for j, col in enumerate(t.columns):
        for c in col.cells:
            c.width = int(usable * w[j] / tot)
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(7.5)
    for row in t.rows:                                   # light borders
        for c in row.cells:
            tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
            for e in ("top", "bottom"):
                x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "BBBBBB"); b.append(x)
            tcPr.append(b)
for p in d.paragraphs:
    for r in p.runs:
        if r.font.size is None: r.font.size = Pt(10)
d.save(sys.argv[1])
